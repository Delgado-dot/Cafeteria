from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


FONT = "Times New Roman"


def set_font(run, size=12, bold=None, italic=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    run.font.color.rgb = None
    run.font.underline = False


def set_para(paragraph, align=None, first_indent=None, hanging=None, before=0, after=0):
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.keep_with_next = False
    fmt.keep_together = False
    if first_indent is not None:
        fmt.first_line_indent = Inches(first_indent)
    if hanging is not None:
        fmt.first_line_indent = Inches(-hanging)
        fmt.left_indent = Inches(hanging)
    if align is not None:
        paragraph.alignment = align
    for run in paragraph.runs:
        set_font(run)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    set_font(run)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def blank_paragraph(p):
    return not p.text.strip()


def is_major_heading(text):
    return bool(re.match(r"^(?:\d+\.\s*)?(Introducci.n|Desarrollo|Conclusi.n|Bibliograf.a|Referencias)$", text, re.I))


def is_subheading(text):
    return bool(re.match(r"^\d+\.\s+.+", text)) and len(text) < 100


def main(src: Path, dst: Path):
    doc = Document(src)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.5)
        section.footer_distance = Inches(0.5)
        header = section.header
        p = header.paragraphs[0]
        p.clear()
        add_page_number(p)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)

    reference_mode = False
    title_phase = True
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if not text:
            continue
        if re.search(r"Bibliograf.a", text, re.I) and len(text) < 50:
            text_new = re.sub(r"Bibliograf.a", "Referencias", p.text, flags=re.I)
            if text_new != p.text:
                for run in p.runs:
                    run.text = ""
                p.runs[0].text = text_new
            text = p.text.strip()
            reference_mode = True

        # The existing first page is retained as the student title page.
        if title_phase and i < 47:
            set_para(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=0)
            if "Computaci" in text:
                for run in p.runs:
                    set_font(run, 12, bold=True)
            continue
        if text.lower() == "contenido":
            title_phase = False
            set_para(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=0)
            for run in p.runs:
                set_font(run, 12, bold=True)
            continue
        title_phase = False

        if reference_mode and not re.search(r"Referencias", text, re.I):
            set_para(p, align=WD_ALIGN_PARAGRAPH.LEFT, hanging=0.5)
        elif is_major_heading(text):
            set_para(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=0)
            for run in p.runs:
                set_font(run, 12, bold=True)
        elif is_subheading(text):
            set_para(p, align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=0)
            for run in p.runs:
                set_font(run, 12, bold=True)
        elif re.match(r"^\d+\.\s+.+\.{3,}\d+$", text):
            set_para(p, align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=0)
        else:
            set_para(p, align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=0.5)

    doc.core_properties.title = "Computación verde"
    doc.core_properties.subject = "Documento formateado según APA 7.ª edición"
    doc.save(dst)


if __name__ == "__main__":
    main(Path(sys.argv[1]), Path(sys.argv[2]))
